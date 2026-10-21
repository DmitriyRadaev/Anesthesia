"""Импорт справочника медикаментов из Excel (перенос artisan-команды `medicament:import`).

Формат файла: первый лист, колонка A — наименование, колонка B — ID группы (1–4).
"""
import json
from pathlib import Path

from django.conf import settings
from django.db import transaction
from openpyxl import load_workbook

DATA_DIR = Path(settings.BASE_DIR) / "data"
DEFAULT_FILE = DATA_DIR / "medicaments.xlsx"
GROUPS_FILE = DATA_DIR / "catalogs" / "medicament_groups.json"


def read_rows(path: Path) -> list[tuple[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for name, group_id, *_ in workbook.active.iter_rows(values_only=True):
        # Как и в исходной системе, чтение прекращается на первой неполной строке.
        if name is None or group_id is None:
            break
        rows.append((str(name).strip(), int(group_id)))
    workbook.close()
    return rows


def load_groups(group_model) -> None:
    for row in json.loads(GROUPS_FILE.read_text(encoding="utf-8")):
        group_model.objects.update_or_create(id=row["id"], defaults={"name": row["name"]})


@transaction.atomic
def import_medicaments(medicament_model, group_model, path: Path = DEFAULT_FILE, replace: bool = False) -> int:
    load_groups(group_model)
    rows = read_rows(path)
    known_groups = set(group_model.objects.values_list("id", flat=True))
    unknown = sorted({g for _, g in rows if g not in known_groups})
    if unknown:
        raise ValueError(f"В файле указаны несуществующие группы медикаментов: {unknown}")
    if replace:
        medicament_model.objects.all().delete()
    created = 0
    for name, group_id in rows:
        _, is_new = medicament_model.objects.update_or_create(name=name, defaults={"group_id": group_id})
        created += int(is_new)
    return created
