"""Отчёт по форме 41.10 «Анестезиологические пособия» (перенос Models/Excel/Report4110 + Report4110Export)."""
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from django.db.models import Prefetch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.catalogs.models import AnesthesiaMethod, TypeOperation
from apps.operations.models import Operation

TYPE_COLUMNS = {
    TypeOperation.URGENT: "anesthesia_urgent",
    TypeOperation.EMERGENCY: "anesthesia_emergency",
    TypeOperation.PLANNED: "anesthesia_planned",
}

ROWS = [
    ("analgosedation", "Аналгоседация"),
    ("epidural", "Эпидуральная анестезия"),
    ("spinal", "Спинальная (субарахноидальная) анестезия"),
    ("spinal_epidural", "Спинально-эпидуральная анестезия"),
    ("tiva", "Тотальная внутривенная анестезия"),
    ("combined_endotracheal", "Комбинированный эндотрахеальный наркоз"),
    ("combined", "Сочетанная анестезия"),
    ("sacral", "Сакральная анестезия"),
    ("intracavitary", "Внутриполосная анестезия"),
    ("total", "Всего"),
]

# Методы анестезии из справочника, по которым строки отчёта считаются автоматически (как в исходной системе).
LOCAL_METHOD_ROWS = {"Эпидуральная": "epidural", "Спинальная (субарахноидальная)": "spinal"}
COUNTED_ROWS = {"epidural", "spinal", "combined", "total"}


@dataclass
class ReportRow:
    key: str
    name: str
    number_line: int
    anesthesia_urgent: int | None = None
    anesthesia_emergency: int | None = None
    anesthesia_planned: int | None = None
    patient_died: int | None = None

    def increment(self, column: str) -> None:
        setattr(self, column, (getattr(self, column) or 0) + 1)


@dataclass
class Report4110:
    date_start: date
    date_end: date
    rows: list[ReportRow] = field(default_factory=list)
    operations_count: int = 0


def build_report_4110(date_start: date, date_end: date) -> Report4110:
    report = Report4110(date_start, date_end, [ReportRow(k, n, i) for i, (k, n) in enumerate(ROWS, start=1)])
    by_key = {row.key: row for row in report.rows}

    def count(row_key: str, type_operation_id: int, died: bool) -> None:
        column = TYPE_COLUMNS.get(type_operation_id)
        for key in (row_key, "total"):
            if column:
                by_key[key].increment(column)
            if died:
                by_key[key].increment("patient_died")

    operations = (
        Operation.objects.filter(status=Operation.Status.FINISHED, operation_date__range=(date_start, date_end),
                                 information_anesthesia__isnull=False)
        .select_related("information_anesthesia", "patient_transferred")
        .prefetch_related(Prefetch("information_anesthesia__anesthesia_methods"))
        .order_by("id")
    )
    for operation in operations:
        report.operations_count += 1
        methods = list(operation.information_anesthesia.anesthesia_methods.all())
        local = [m.name for m in methods if m.type == AnesthesiaMethod.Kind.LOCAL]
        general = [m.name for m in methods if m.type == AnesthesiaMethod.Kind.GENERAL]
        transferred = getattr(operation, "patient_transferred", None)
        died = bool(transferred and transferred.is_died)

        for name in local:
            if name in LOCAL_METHOD_ROWS:
                count(LOCAL_METHOD_ROWS[name], operation.type_operation_id, died)
        if local and general:
            count("combined", operation.type_operation_id, died)
    return report


def _fmt(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def report_filename(report: Report4110) -> str:
    return f"41.10 ({_fmt(report.date_start)}-{_fmt(report.date_end)}).xlsx"


def render_report_4110_xlsx(report: Report4110) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "41.10"
    thin, thick = Side(style="thin"), Side(style="medium")
    header_fill = PatternFill("solid", start_color="EAEDF6")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Протокол 41.10 за период: {_fmt(report.date_start)} - {_fmt(report.date_end)}"
    ws["A1"].font = Font(italic=True, size=9)
    ws["A1"].alignment = Alignment(horizontal="right")

    ws.merge_cells("A2:A3"); ws["A2"] = "Вид анестезии"
    ws.merge_cells("B2:B3"); ws["B2"] = "№ строки"
    ws.merge_cells("C2:E2"); ws["C2"] = "Проведено анестезий, ед"
    ws.merge_cells("F2:F3"); ws["F2"] = "Умерло, чел"
    ws["C3"], ws["D3"], ws["E3"] = "Неотложных", "Экстренных", "Плановых"
    for col, number in zip("ABCDEF", range(1, 7)):
        ws[f"{col}4"] = number

    for row_index, row in enumerate(report.rows, start=5):
        values = [row.name, row.number_line, row.anesthesia_urgent, row.anesthesia_emergency,
                  row.anesthesia_planned, row.patient_died]
        for col_index, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col_index, value=value)

    last = 4 + len(report.rows)
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=6):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = center
            if cell.row <= 4:
                cell.fill = header_fill
                cell.font = Font(bold=cell.row <= 3, size=10 if cell.row <= 3 else 9)
            if cell.column == 1 and cell.row > 4:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for cell in ws[last]:
        cell.font = Font(bold=True)
        cell.border = Border(left=thin, right=thin, top=thick, bottom=thick)

    ws.column_dimensions["A"].width = 45
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
